"""Prepare unlabelled Singlish candidates without changing source text.

Run: python src/prepare_singlish.py

Exclusion precedence is null/empty, whitespace-only, marker-only, then exact
duplicates. Each removed row receives exactly one removal reason. Sample IDs
use the original Excel data-row position, so excluded rows leave ID gaps.
Review reasons are flags for humans, never toxicity or verified language labels.
Existing identical outputs are reusable; different outputs are never overwritten.
"""

import argparse
from collections import Counter
import csv
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import unicodedata

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = PROJECT_ROOT / "data/raw/singlish/singlish_raw.xlsx"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "data/interim/singlish_annotation"
AUDITED_SOURCE_SHA256 = (
    "e01d06123579c4a45a446766b4907b5681390a85c1743426c23b0a0a1e99bab8"
)
AUDITED_PROVENANCE_ROWS = (35526, 40525)
SHEET_NAME = "Comments"
TEXT_COLUMN = "SINGLISH"
CANDIDATE_COLUMNS = ("sample_id", "text", "source_row", "review_status", "review_reason")
REMOVED_COLUMNS = ("source_row", "text", "removal_reason")
REMOVAL_REASONS = ("null_empty", "whitespace_only", "marker_only", "exact_duplicate")
REVIEW_REASONS = (
    "partial_unknown_marker",
    "possible_phone_number",
    "non_latin_script",
    "single_word",
    "leading_quote",
    "timestamp_like",
    "uncertain_provenance",
)
UNKNOWN_MARKER = re.compile(r"\[\s*unkown\s*\]", re.IGNORECASE)
PHONE_LIKE = re.compile(r"(?<!\w)\+?\d(?:[\d ()-]{5,}\d)(?!\w)")
TIMESTAMP_LIKE = re.compile(r"(?<!\d)(?:\d{1,2}:)?\d{1,2}:\d{2}(?:[,.]\d{1,3})?(?!\d)")
WORD = re.compile(r"[^\W\d_]+", re.UNICODE)


class PreparationError(ValueError):
    """A preparation failure with a diagnostic that contains no comment text."""


def sha256_file(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def is_marker_only(text):
    """Ignore separators for detection only; never rewrite the stored text."""
    if not UNKNOWN_MARKER.search(text):
        return False
    remainder = UNKNOWN_MARKER.sub("", text)
    return all(char.isspace() or unicodedata.category(char).startswith("P")
               for char in remainder)


def review_reasons(text, source_row, provenance_rows=None):
    """Return consistently ordered, non-exclusive human-review flags."""
    reasons = []
    if UNKNOWN_MARKER.search(text):
        reasons.append("partial_unknown_marker")
    if any(7 <= sum(char.isdigit() for char in match.group()) <= 15
           for match in PHONE_LIKE.finditer(text)):
        reasons.append("possible_phone_number")
    if any(char.isalpha() and "LATIN" not in unicodedata.name(char, "")
           for char in text):
        reasons.append("non_latin_script")
    if len(WORD.findall(text)) == 1:
        reasons.append("single_word")
    if text.lstrip().startswith(('"', "'", "\u2018", "\u201c")):
        reasons.append("leading_quote")
    if TIMESTAMP_LIKE.search(text):
        reasons.append("timestamp_like")
    if provenance_rows and provenance_rows[0] <= source_row <= provenance_rows[1]:
        reasons.append("uncertain_provenance")
    return reasons


def prepare_records(records, provenance_rows=None):
    """Partition (Excel row, original value) pairs; all text stays unchanged."""
    candidates, removed = [], []
    seen = set()
    raw_seen = set()
    exclusions = Counter({reason: 0 for reason in REMOVAL_REASONS})
    reviews = Counter({reason: 0 for reason in REVIEW_REASONS})
    original_rows = raw_duplicates = 0
    previous_row = 1
    for source_row, text in records:
        if not isinstance(source_row, int) or source_row <= previous_row:
            raise PreparationError("Source rows must be ascending Excel row numbers, starting at 2.")
        previous_row = source_row
        original_rows += 1
        if text is not None and not isinstance(text, str):
            raise PreparationError(f"Non-text comment at Excel row {source_row}; manual review required.")
        if isinstance(text, str) and text != "":
            raw_duplicates += text in raw_seen
            raw_seen.add(text)

        reason = None
        if text is None or text == "":
            reason = "null_empty"
        elif not text.strip():
            reason = "whitespace_only"
        elif is_marker_only(text):
            reason = "marker_only"
        elif text in seen:
            reason = "exact_duplicate"

        if reason:
            exclusions[reason] += 1
            removed.append({"source_row": source_row, "text": text,
                            "removal_reason": reason})
            continue
        seen.add(text)
        reasons = review_reasons(text, source_row, provenance_rows)
        reviews.update(reasons)
        candidates.append({
            "sample_id": f"SI{source_row - 1:05d}",
            "text": text,
            "source_row": source_row,
            "review_status": "needs_review" if reasons else "pending",
            "review_reason": ";".join(reasons),
        })

    needs_review = sum(row["review_status"] == "needs_review" for row in candidates)
    summary = {
        "original_rows": original_rows,
        "exclusions": dict(exclusions),
        "raw_exact_duplicate_occurrences": raw_duplicates,
        "candidate_rows_remaining": len(candidates),
        "pending_rows": len(candidates) - needs_review,
        "needs_review_rows": needs_review,
        "review_reason_counts": dict(reviews),
    }
    assert original_rows == len(candidates) + len(removed)
    assert len(removed) == sum(exclusions.values())
    return candidates, removed, summary


def read_records(source_bytes):
    """Read an immutable workbook snapshot, without evaluating formulas."""
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(source_bytes), read_only=True, data_only=False, keep_links=False
        )
    except Exception as error:
        # Parser exception messages may contain cell contents; expose only type.
        raise PreparationError(f"Could not open source workbook ({type(error).__name__}).") from None
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise PreparationError(f"Missing worksheet: {SHEET_NAME}.")
        rows = workbook[SHEET_NAME].iter_rows()
        header = [cell.value for cell in next(rows, ())]
        if header.count(TEXT_COLUMN) != 1:
            raise PreparationError(f"Expected exactly one {TEXT_COLUMN} column in the first row.")
        column = header.index(TEXT_COLUMN)
        records = []
        # Some source rows include blank cells through AZ and have no dimension metadata.
        for source_row, cells in enumerate(rows, start=2):
            cell = cells[column] if column < len(cells) else None
            if cell is not None and cell.data_type in ("f", "e"):
                raise PreparationError(f"Formula/error comment at Excel row {source_row}; review required.")
            records.append((source_row, cell.value if cell is not None else None))
        return records
    except PreparationError:
        raise
    except Exception as error:
        raise PreparationError(f"Could not read source cells ({type(error).__name__}).") from None
    finally:
        workbook.close()


def csv_bytes(rows, columns):
    buffer = io.StringIO(newline="")
    writer = csv.DictWriter(buffer, fieldnames=columns)
    writer.writeheader()
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8-sig")


def prepare_workbook(source_path=DEFAULT_SOURCE, output_dir=DEFAULT_OUTPUT_DIR):
    source_path, output_dir = Path(source_path).resolve(), Path(output_dir).resolve()
    source_bytes = source_path.read_bytes()
    before = hashlib.sha256(source_bytes).hexdigest()
    # The row-range finding belongs to this exact audited source, not arbitrary workbooks.
    provenance_rows = AUDITED_PROVENANCE_ROWS if before == AUDITED_SOURCE_SHA256 else None
    if source_path == DEFAULT_SOURCE.resolve() and provenance_rows is None:
        raise PreparationError("The raw source differs from the audited workbook; re-audit before preparation.")
    candidates, removed, summary = prepare_records(read_records(source_bytes), provenance_rows)
    after = sha256_file(source_path)
    if before != after:
        raise PreparationError("Source SHA-256 changed during processing; no outputs published.")
    summary.update({
        "source_file": source_path.as_posix(),
        "source_sheet": SHEET_NAME,
        "text_column": TEXT_COLUMN,
        "source_sha256_before": before,
        "source_sha256_after": after,
        "source_unchanged": True,
        "exclusion_precedence": list(REMOVAL_REASONS),
        "review_reason_delimiter": ";",
        "sample_id_rule": "SI + (original Excel row - 1), zero-padded to at least 5 digits",
        "provenance_review_excel_rows": provenance_rows,
        "language_policy": "No language classification. Non-Latin letters are flagged; pending does not verify language eligibility.",
    })
    outputs = {
        "singlish_cleaned_unlabelled.csv": csv_bytes(candidates, CANDIDATE_COLUMNS),
        "singlish_cleaning_removed.csv": csv_bytes(removed, REMOVED_COLUMNS),
        "singlish_preparation_summary.json": (json.dumps(summary, indent=2) + "\n").encode("utf-8"),
    }
    # Check every destination before creating any files. Preserve human edits on reruns.
    for filename, content in outputs.items():
        destination = output_dir / filename
        if destination.exists() and destination.read_bytes() != content:
            raise PreparationError(f"Refusing to overwrite different existing output: {filename}.")
    output_dir.mkdir(parents=True, exist_ok=True)
    created = []
    try:
        for filename, content in outputs.items():
            destination = output_dir / filename
            if not destination.exists():
                with destination.open("xb") as output:
                    created.append(destination)
                    output.write(content)
        if sha256_file(source_path) != before:
            raise PreparationError("Source SHA-256 changed before completion; new outputs withdrawn.")
    except Exception:
        for destination in created:
            destination.unlink()
        raise
    return summary


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR,
                        help="Destination for the two interim CSVs and JSON summary.")
    args = parser.parse_args(argv)
    try:
        summary = prepare_workbook(output_dir=args.output_dir)
    except (OSError, PreparationError) as error:
        print(f"Preparation failed: {error}", file=sys.stderr)
        return 1
    # Counts, paths and hashes only: never print source text or phone-like matches.
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
